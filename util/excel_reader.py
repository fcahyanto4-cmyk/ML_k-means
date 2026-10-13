from openpyxl import load_workbook
from datetime import datetime, timedelta

def read_excel(file_path):
    wb = load_workbook(file_path)
    sheet = wb.active

    data = []

    # Ambil header dulu
    header = [cell.value for cell in sheet[1]]
    print(f"Header: {header}")  # Debug header

    # Buat mapping nama kolom ke indeks
    kolom = {name.lower(): idx for idx, name in enumerate(header)}

    for row in sheet.iter_rows(min_row=2, values_only=True):
        try:
            nama = row[kolom["nama wajib pajak"]]
            npwp = row[kolom["npwp"]]
            jenis = row[kolom["jenis pajak"]]
            tanggal = row[kolom["tanggal bayar"]]
            jumlah = row[kolom["jumlah bayar (rp)"]]  # Update sesuai dengan nama kolom yang benar

            # Cek apakah 'jumlah' adalah angka, jika tidak, skip
            if not isinstance(jumlah, (int, float)):
                print(f"Jumlah bayar invalid: {jumlah}")
                continue

            # Pastikan format tanggal sesuai
            if isinstance(tanggal, (int, float)):
               tanggal = datetime.fromordinal(datetime(1900, 1, 1).toordinal() + int(tanggal) - 2)

            data.append({
                "nama": nama,
                "npwp": npwp,
                "jenis": jenis,
                "tanggal": tanggal,
                "jumlah": jumlah
            })
        except Exception as e:
            print(f"Baris dilewati karena error: {e}")

    print(f"Data yang berhasil dibaca: {data}")  # Debug data yang berhasil dibaca
    return data
            
# Fungsi untuk menghitung denda berdasarkan keterlambatan
def hitung_denda(tanggal_bayar, jumlah_bayar):
    # Tentukan tanggal jatuh tempo (misalnya tanggal 15 setiap bulan)
    tanggal_jatuh_tempo = datetime(tanggal_bayar.year, tanggal_bayar.month, 15)
    
    if tanggal_bayar > tanggal_jatuh_tempo:
        # Hitung jumlah hari keterlambatan
        keterlambatan = (tanggal_bayar - tanggal_jatuh_tempo).days
        
        # Aturan denda 1% per hari keterlambatan
        denda = jumlah_bayar * 0.01 * keterlambatan
        return denda
    else:
        return 0

# Fungsi untuk menganalisis tren dan menghitung total, rata-rata, dan denda
# Pastikan bahwa total_pembayaran, total_denda, dan lainnya adalah angka
def analisis_tren(data):
    hasil = {}
    for d in data:
        nama = d['Nama Wajib Pajak']
        jumlah = float(d['Jumlah Bayar (Rp)'])
        denda = float(d['Denda'])

        if nama not in hasil:
            hasil[nama] = {
                'total_pembayaran': 0,
                'total_denda': 0,
                'jumlah_transaksi': 0
            }

        hasil[nama]['total_pembayaran'] += jumlah
        hasil[nama]['total_denda'] += denda
        hasil[nama]['jumlah_transaksi'] += 1

    # Tambahkan rata-rata pembayaran
    for nama, info in hasil.items():
        if info['jumlah_transaksi'] > 0:
            info['rata_rata_pembayaran'] = info['total_pembayaran'] / info['jumlah_transaksi']
        else:
            info['rata_rata_pembayaran'] = 0

        # Pastikan tidak ada None atau nilai tidak valid
        info['total_pembayaran'] = info['total_pembayaran'] or 0
        info['total_denda'] = info['total_denda'] or 0
        info['rata_rata_pembayaran'] = info['rata_rata_pembayaran'] or 0

    return hasil
